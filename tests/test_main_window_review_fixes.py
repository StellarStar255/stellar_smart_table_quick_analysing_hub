"""主窗口第二轮审查修复的回归测试：

- 预览框 400ms 延迟写回：换 sheet / 新建 / 关窗前先落盘，模型整表重置后作废旧坐标
- .tsv/.txt 走文本保存，不再被当成 xlsx 写成 zip
- 多 sheet 工作簿"保存为 CSV"按导出处理：当前文件、修改标记、其他 sheet 原样保留
- 粘贴超出末行用可撤销的追加行，不清空撤销栈
- 整列/全选的批量操作按选区范围推导，不逐格枚举 QModelIndex
- 「复制列名」偏好真正持久化
- 自动保存：输入中不弹窗；有损文件拒绝一次后停用该文件的自动保存
- 拖放打开延后到下一轮事件循环；关窗先关分析窗口；工作簿句柄换文件时关闭
- 语言菜单勾选靠引用刷新；重启命令用 argv[0]
- 图片面板整表替换后只同步一次

需要 QApplication（widgets）；CI 里用 QT_QPA_PLATFORM=offscreen 无头运行。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
import pytest
from PyQt6.QtCore import (
    Qt, QSettings, QItemSelection, QItemSelectionModel, QMimeData, QUrl, QPointF,
)
from PyQt6.QtGui import QDropEvent
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import (
    QApplication, QMessageBox, QFileDialog, QAbstractItemView, QMenu,
)

_app = QApplication.instance() or QApplication([])

from qtui import file_io
from qtui import main_window as mw_mod
from qtui.main_window import MainWindow, tr

SELECT = QItemSelectionModel.SelectionFlag.Select
CLEAR_SELECT = QItemSelectionModel.SelectionFlag.ClearAndSelect


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    ini = str(tmp_path / "settings.ini")
    monkeypatch.setattr(mw_mod, "_make_settings",
                        lambda: QSettings(ini, QSettings.Format.IniFormat))


@pytest.fixture
def win():
    w = MainWindow()
    w.model.set_dataframe(pd.DataFrame({'A': [1.0, 2.0, 3.0], 'B': [4.0, 5.0, 6.0]}))
    w.model.modified = False
    yield w
    w._analysis_win = None
    w.model.modified = False
    w.close()


def _select(win, top, left, bottom, right, clear=True):
    win.table.selectionModel().select(
        QItemSelection(win.model.index(top, left), win.model.index(bottom, right)),
        CLEAR_SELECT if clear else SELECT)


def _two_memory_sheets(win):
    """当前表变成 S1，再加一个只在内存里的 S2。"""
    win.sheet_names = ["S1", "S2"]
    win.current_sheet = "S1"
    win._cache_sheet("S1", win.model.df, pin=True)
    win._cache_sheet("S2", pd.DataFrame({'X': [10.0, 20.0, 30.0]}), pin=True)


def _type_in_preview(win, row, col, text):
    win.table.setCurrentIndex(win.model.index(row, col))
    win.cell_preview_text.setPlainText(text)
    assert win._preview_save_timer.isActive()


# ================= 预览框延迟写回 =================

class TestPreviewFlush:
    def test_switch_sheet_writes_pending_edit_to_old_sheet(self, win):
        _two_memory_sheets(win)
        _type_in_preview(win, 1, 0, "typed")
        win.switch_sheet("S2")
        assert not win._preview_save_timer.isActive()
        assert win._preview_cell is None
        assert str(win._sheet_cache["S1"].iat[0, 0]) == "typed"   # 写回了原 sheet
        assert win.model.df.iat[0, 0] == 10.0                      # 新 sheet 没被覆盖
        assert win.cell_preview_text.toPlainText() == ""

    def test_new_file_flushes_before_reset(self, win, monkeypatch):
        old_df = win.model.df
        _type_in_preview(win, 1, 1, "late")
        monkeypatch.setattr(QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.Discard)
        win.new_file()
        assert str(old_df.iat[0, 1]) == "late"
        assert win.model.df.isna().all().all()     # 新表干净

    def test_model_reset_cancels_stale_preview_write(self, win):
        _type_in_preview(win, 1, 0, "stale")
        win.model.set_dataframe(pd.DataFrame({'Z': [0.0, 0.0]}))
        assert not win._preview_save_timer.isActive()
        assert win._preview_cell is None
        win._save_cell_preview()                    # 即便被触发也无处可写
        assert win.model.df.iat[0, 0] == 0.0

    def test_filter_reapply_flushes(self, win):
        _type_in_preview(win, 1, 0, "9")
        win.active_filters = [{"col": "B", "condition": "大于", "value": "4"}]
        win._reapply_filters()
        assert str(win.original_df.iat[0, 0]) in ("9", "9.0")
        assert not win._preview_save_timer.isActive()
        win.clear_all_filters()

    def test_close_event_flushes_pending_edit(self, win, monkeypatch):
        _type_in_preview(win, 1, 0, "bye")
        asked = []
        monkeypatch.setattr(QMessageBox, "question",
                            lambda *a, **k: asked.append(1) or QMessageBox.StandardButton.Cancel)
        assert win.close() is False
        assert asked                                # 写回产生修改 → 提示保存而非静默丢失
        assert str(win.model.df.iat[0, 0]) == "bye"


# ================= 文本格式 =================

class TestTextFormats:
    def test_helper_recognises_text_extensions(self):
        assert mw_mod._is_text_format("a.tsv")
        assert mw_mod._is_text_format("A.TXT")
        assert mw_mod._is_text_format("x.csv")
        assert not mw_mod._is_text_format("x.xlsx")

    def test_save_tsv_writes_text_not_zip(self, win, tmp_path):
        path = str(tmp_path / "data.tsv")
        win.model.modified = True
        assert win._do_save(path) is True
        with open(path, "rb") as fh:
            assert fh.read(2) != b"PK"              # 不是 xlsx 的 zip
        first = open(path, encoding="utf-8-sig").read().splitlines()[0]
        assert first.startswith("A")
        assert win.current_file == path and win.model.modified is False

    @pytest.mark.skipif("sep" not in file_io.save_csv.__code__.co_varnames,
                        reason="file_io.save_csv 尚不支持按后缀推断分隔符")
    def test_save_tsv_is_tab_separated(self, win, tmp_path):
        path = str(tmp_path / "data.tsv")
        win.model.modified = True
        assert win._do_save(path) is True
        first = open(path, encoding="utf-8-sig").read().splitlines()[0]
        assert first == "A\tB"
        back = file_io.read_csv_any_encoding(path)
        assert list(back.columns) == ["A", "B"] and len(back) == 3

    def test_csv_save_message_mentions_dropped_colors(self, win, tmp_path):
        win.model.cell_colors[(0, 0)] = "#a04040"
        win.model.modified = True
        assert win._do_save(str(tmp_path / "c.csv")) is True
        msg = win.statusBar().currentMessage()
        assert tr("已保存: {}").format(str(tmp_path / "c.csv")) in msg
        assert tr("提示：CSV 格式不保存背景色，用 xlsx 可保留") in msg

    def test_openable_extensions_include_txt(self):
        assert ".txt" in mw_mod._OPENABLE_EXTS and ".tsv" in mw_mod._OPENABLE_EXTS


# ================= 多 sheet 工作簿另存为 CSV = 导出 =================

class TestMultiSheetTextExport:
    def _workbook(self, win, tmp_path):
        xlsx = str(tmp_path / "book.xlsx")
        with pd.ExcelWriter(xlsx) as xw:
            pd.DataFrame({'A': [1, 2]}).to_excel(xw, sheet_name="S1", index=False)
            pd.DataFrame({'B': [3, 4]}).to_excel(xw, sheet_name="S2", index=False)
        win.current_file = xlsx
        win._excel_file, win.sheet_names = file_io.load_workbook_lazy(xlsx)
        win.current_sheet = "S1"
        win.model.set_dataframe(file_io.read_sheet(win._excel_file, "S1"))
        win.model.setData(win.model.index(1, 0), "99")
        return xlsx

    def test_save_as_csv_keeps_workbook_open(self, win, tmp_path, monkeypatch):
        xlsx = self._workbook(win, tmp_path)
        csv_path = str(tmp_path / "out.csv")
        monkeypatch.setattr(QFileDialog, "getSaveFileName",
                            staticmethod(lambda *a, **k: (csv_path, "")))
        monkeypatch.setattr(QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.Yes)
        assert win._save_as(switch_to=True) is True
        assert os.path.exists(csv_path)
        assert win.current_file == xlsx             # 没有换成 .csv
        assert win.sheet_names == ["S1", "S2"]
        assert win.model.modified is True           # 工作簿本身仍待保存
        assert "out.csv" in win.statusBar().currentMessage()
        assert os.path.basename(xlsx) in win.windowTitle()

    def test_discard_prompt_does_not_treat_export_as_saved(self, win, tmp_path, monkeypatch):
        _two_memory_sheets(win)                     # 未命名的多 sheet 工作簿
        win.model.modified = True
        csv_path = str(tmp_path / "out.csv")
        monkeypatch.setattr(QFileDialog, "getSaveFileName",
                            staticmethod(lambda *a, **k: (csv_path, "")))
        answers = iter([QMessageBox.StandardButton.Save, QMessageBox.StandardButton.Yes])
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: next(answers))
        assert win._check_save_before_discard() is False
        assert os.path.exists(csv_path)
        assert win.model.modified is True and win.current_file is None


# ================= 粘贴超出末行 =================

class TestPasteAppendRows:
    def test_paste_beyond_last_row_keeps_undo_history(self, win):
        win.model.setData(win.model.index(1, 0), "7")
        win.table.setCurrentIndex(win.model.index(3, 0))       # 最后一个数据行
        QApplication.clipboard().setText("x\ny\nz")
        win.paste_selection()
        assert len(win.model.df) == 5
        assert [str(v) for v in win.model.df.iloc[2:5, 0]] == ["x", "y", "z"]
        while win.model.undo():
            pass
        assert len(win.model.df) == 3                          # 追加的行撤掉了
        assert win.model.df.iat[0, 0] == 1.0                   # 粘贴前的编辑也能撤回


# ================= 选区按范围推导 =================

class TestRangeBasedSelection:
    def test_delete_rows_from_whole_column_selection(self, win, monkeypatch):
        win.table.selectColumn(0)
        monkeypatch.setattr(QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.Yes)
        win.delete_selected_rows()
        assert len(win.model.df) == 0

    def test_delete_columns_from_selection(self, win, monkeypatch):
        win.table.selectColumn(1)
        monkeypatch.setattr(QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.Yes)
        win.delete_selected_columns()
        assert list(win.model.df.columns) == ["A"]

    def test_apply_function_over_column(self, win, monkeypatch):
        win.table.selectColumn(0)
        monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
        win.apply_function("sum")
        assert win.statusBar().currentMessage().endswith(": 6.0")
        win.apply_function("count")
        assert win.statusBar().currentMessage().endswith(": 3")

    def test_copy_disjoint_ranges(self, win):
        win.copy_headers_cb.setChecked(False)
        _select(win, 1, 0, 1, 0)
        _select(win, 3, 1, 3, 1, clear=False)
        win.copy_selection()
        assert QApplication.clipboard().text() == "1.0\t\n\t6.0"

    def test_copy_disjoint_ranges_carries_formulas(self, win):
        win.model.setData(win.model.index(3, 1), "=A4*2")
        win.copy_headers_cb.setChecked(False)
        _select(win, 1, 0, 1, 0)
        _select(win, 3, 1, 3, 1, clear=False)
        win.copy_selection()
        assert win._formula_clipboard["cells"] == {
            (1, 1): {"text": "=A4*2", "src": (3, 1)}}

    def test_selection_frame(self, win):
        win.table.selectColumn(1)
        frame, names = win.selection_frame()
        assert names == ["B"] and list(frame["B"]) == [4.0, 5.0, 6.0]

    def test_color_and_transposed_copy(self, win):
        win.table.selectColumn(0)
        win._set_selection_color("#a04040")
        assert win.model.cell_colors.get((0, 0)) == "#a04040"
        assert win.model.cell_colors.get((2, 0)) == "#a04040"
        win.copy_selection_transposed()
        assert QApplication.clipboard().text() == "A\t1.0\t2.0\t3.0"

    def test_ctrl_enter_fills_selection_from_ranges(self, win):
        win.table.setCurrentIndex(win.model.index(1, 0))
        win.table.edit(win.model.index(1, 0))
        editor = win._cell_delegate.active_editor
        editor.setText("5")
        _select(win, 1, 0, 3, 0)       # 打开编辑器后再拉选区（edit 会重置选区）
        win.table.commit_and_fill(editor)
        assert list(win.model.df["A"].astype(float)) == [5.0, 5.0, 5.0]


# ================= 「复制列名」偏好持久化 =================

class TestCopyHeadersPreference:
    def test_toggle_persists_across_windows(self):
        w1 = MainWindow()
        try:
            assert not w1.copy_headers_cb.isChecked()
            w1.copy_headers_cb.setChecked(True)
            w1._settings.sync()
        finally:
            w1.model.modified = False
            w1.close()
        w2 = MainWindow()
        try:
            assert w2.copy_headers_cb.isChecked()
        finally:
            w2.model.modified = False
            w2.close()


# ================= 自动保存 =================

class TestAutoSave:
    def test_skipped_while_editing(self, win, monkeypatch, tmp_path):
        win.current_file = str(tmp_path / "f.csv")
        win.model.modified = True
        saves = []
        monkeypatch.setattr(win, "_do_save", lambda *a, **k: saves.append(a) or True)
        win.auto_save_cb.setChecked(True)
        win.table.setCurrentIndex(win.model.index(1, 0))
        win.table.edit(win.model.index(1, 0))
        assert win.table.state() == QAbstractItemView.State.EditingState
        win._do_auto_save()
        assert saves == []
        assert win._auto_save_timer.isActive()      # 稍后重试，而不是放弃
        win.table.setCurrentIndex(win.model.index(2, 0))
        win.table.closePersistentEditor(win.model.index(1, 0))

    def test_declined_lossy_confirm_disables_auto_save_for_file(self, win, monkeypatch, tmp_path):
        path = str(tmp_path / "book.xlsx")
        pd.DataFrame({'A': [1]}).to_excel(path, index=False)
        win.current_file = path
        win._excel_file, win.sheet_names = file_io.load_workbook_lazy(path)
        win.current_sheet = win.sheet_names[0]
        win.model.modified = True
        asked, saves = [], []

        def decline(src):
            asked.append(src)
            return False
        monkeypatch.setattr(win, "_confirm_lossy_save", decline)
        monkeypatch.setattr(win, "_do_save", lambda *a, **k: saves.append(a) or True)
        win.auto_save_cb.setChecked(True)           # 勾选时问一次
        win._do_auto_save()
        win._do_auto_save()
        assert len(asked) == 1 and saves == []      # 不再反复弹窗、不保存
        assert path in win._auto_save_declined
        win._mark_modified()
        assert not win._auto_save_timer.isActive()
        assert os.path.basename(path) in win.statusBar().currentMessage()
        # 重新勾选 = 再给一次机会
        win.auto_save_cb.setChecked(False)
        win.auto_save_cb.setChecked(True)
        assert len(asked) == 2

    def test_accepted_lossy_confirm_saves(self, win, monkeypatch, tmp_path):
        path = str(tmp_path / "book.xlsx")
        pd.DataFrame({'A': [1]}).to_excel(path, index=False)
        win.current_file = path
        win._excel_file, win.sheet_names = file_io.load_workbook_lazy(path)
        win.current_sheet = win.sheet_names[0]
        win.model.modified = True
        saves = []
        monkeypatch.setattr(win, "_confirm_lossy_save", lambda src: True)
        monkeypatch.setattr(win, "_do_save", lambda *a, **k: saves.append(a) or True)
        win.auto_save_cb.setChecked(True)
        win._do_auto_save()
        assert saves == [(path,)]


# ================= 拖放 / 关窗 / 句柄 =================

class _FakeExcelFile:
    sheet_names = ["S"]

    def __init__(self):
        self.closed = False

    def close(self):
        self.closed = True


class TestDropAndClose:
    def test_drop_defers_load_to_event_loop(self, win, monkeypatch, tmp_path):
        path = str(tmp_path / "d.txt")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("A\tB\n1\t2\n")
        loaded = []
        monkeypatch.setattr(win, "load_file", lambda p: loaded.append(p))
        mime = QMimeData()
        mime.setUrls([QUrl.fromLocalFile(path)])
        event = QDropEvent(QPointF(10, 10), Qt.DropAction.CopyAction, mime,
                           Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
        win.dropEvent(event)
        assert loaded == []                         # 拖放事件里不同步加载
        QTest.qWait(50)
        assert loaded == [path]

    def test_refusing_analysis_window_blocks_close(self, win):
        class FakeAnalysis:
            calls = 0

            def isVisible(self):
                return True

            def close(self):
                self.calls += 1
                return False
        fake = FakeAnalysis()
        win._analysis_win = fake
        assert win.close() is False
        assert fake.calls == 1

    def test_release_closes_previous_handle(self, win):
        fake = _FakeExcelFile()
        win._excel_file = fake
        win._release_excel_file()
        assert fake.closed and win._excel_file is None

    def test_load_file_closes_previous_handle(self, win, tmp_path):
        fake = _FakeExcelFile()
        win._excel_file = fake
        path = str(tmp_path / "n.xlsx")
        pd.DataFrame({'A': [1]}).to_excel(path, index=False)
        win.load_file(path)
        _app.processEvents()
        assert fake.closed
        assert win._excel_file is not fake
        assert win._coord_marker_cache.get(path) is False   # pandas 写的文件没有标记
        win._release_excel_file()


# ================= 语言菜单 / 重启命令 / 菜单辅助 =================

class TestMenusAndRestart:
    def test_language_checks_refresh_via_refs(self, win):
        from qtui import i18n
        win._build_menu_refresh_language_checks(i18n.LANG_EN)
        assert win._lang_actions[i18n.LANG_EN].isChecked()
        assert not win._lang_actions[i18n.LANG_ZH].isChecked()
        win._build_menu_refresh_language_checks(i18n.LANG_ZH)
        assert win._lang_actions[i18n.LANG_ZH].isChecked()

    def test_restart_command(self, monkeypatch, tmp_path):
        script = tmp_path / "entry.py"
        script.write_text("", encoding="utf-8")
        monkeypatch.setattr(sys, "argv", [str(script)])
        monkeypatch.delattr(sys, "frozen", raising=False)
        assert MainWindow._restart_command() == [sys.executable, str(script)]
        monkeypatch.setattr(sys, "argv", ["-c"])
        cmd = MainWindow._restart_command()
        assert cmd[1].endswith("smart_table_quick_analysing_hub.py")
        monkeypatch.setattr(sys, "frozen", True, raising=False)
        assert MainWindow._restart_command() == [sys.executable]

    def test_row_and_column_action_helpers(self, win):
        menu = QMenu()
        win._add_row_actions(menu, 2)
        texts = [a.text() for a in menu.actions()]
        assert texts == [tr("向上插入一行"), tr("向下插入一行"),
                         tr("删除选中行"), tr("将此行设为表头")]
        menu = QMenu()
        win._add_row_actions(menu, 0)
        assert tr("将此行设为表头") not in [a.text() for a in menu.actions()]
        header = QMenu()
        win._add_column_actions(header, 1, header_menu=True)
        header_texts = [a.text() for a in header.actions()]
        assert tr("删除此列") in header_texts and tr("删除选中列") not in header_texts
        cell = QMenu()
        win._add_column_actions(cell, 1)
        cell_texts = [a.text() for a in cell.actions()]
        assert tr("删除选中列") in cell_texts and tr("删除此列") not in cell_texts


# ================= 图片面板 / 后台读 sheet =================

class TestSheetSwitching:
    def test_switch_sheet_syncs_image_panel_once(self, win, monkeypatch):
        _two_memory_sheets(win)
        calls = []
        orig = win.image_panel.set_context
        monkeypatch.setattr(win.image_panel, "set_context",
                            lambda *a, **k: calls.append(a) or orig(*a, **k))
        win.switch_sheet("S2")
        assert len(calls) == 1

    def test_switch_sheet_reads_formulas_and_caches_marker(self, win, tmp_path):
        import openpyxl
        path = str(tmp_path / "f.xlsx")
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "S1"
        ws1.append(["A"])
        ws1.append([1])
        ws2 = wb.create_sheet("S2")
        ws2.append(["A", "B"])
        ws2.append([2, "=A2*2"])
        wb.save(path)
        win.current_file = path
        win._excel_file, win.sheet_names = file_io.load_workbook_lazy(path)
        win.current_sheet = "S1"
        win.model.set_dataframe(file_io.read_sheet(win._excel_file, "S1"))
        win.model.modified = False
        win.switch_sheet("S2")
        assert win.current_sheet == "S2"
        assert any(f.upper().startswith("=A2*2") for f in win.model.formulas.values())
        assert win._coord_marker_cache == {path: False}
        win._release_excel_file()

    def test_switch_sheet_failure_keeps_current(self, win, monkeypatch):
        _two_memory_sheets(win)
        win.current_file = "/nonexistent/book.xlsx"
        win._excel_file = _FakeExcelFile()          # 声称有 S 这个 sheet，但读不出来
        win.sheet_names = ["S1", "S"]
        errors = []
        monkeypatch.setattr(win, "_show_error", lambda *a, **k: errors.append(a))
        monkeypatch.setattr(file_io, "read_sheet",
                            lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
        win.switch_sheet("S")
        assert errors and win.current_sheet == "S1"
        win._excel_file = None


# ================= 批量写入（set_cells）：一次撤销 / 一次 dataChanged / 筛选镜像 =================

def _count_data_changes(model):
    """统计携带 Display/Edit 角色的 dataChanged 次数（忽略行高亮等背景色刷新）。"""
    hits = []

    def on_changed(tl, br, roles=None):
        roles = list(roles or [])
        if (not roles or Qt.ItemDataRole.DisplayRole in roles
                or Qt.ItemDataRole.EditRole in roles):
            hits.append((tl.row(), tl.column(), br.row(), br.column()))
    model.dataChanged.connect(on_changed)
    return hits


class TestBatchCellWrites:
    @pytest.fixture
    def big(self, win):
        win.model.set_dataframe(pd.DataFrame({
            'A': [1.0, 2.0, 3.0, 4.0, 5.0, 6.0],
            'B': [10.0, 20.0, 30.0, 40.0, 50.0, 60.0],
            'C': ['p', 'q', 'r', 's', 't', 'u'],
        }))
        win.model.modified = False
        return win

    def test_paste_block_single_undo_and_single_data_changed(self, big):
        win = big
        before = win.model.df.copy()
        hits = _count_data_changes(win.model)
        win.table.setCurrentIndex(win.model.index(1, 0))
        QApplication.clipboard().setText(
            "\n".join("\t".join(f"{r}{c}" for c in "xyz") for r in range(5)))
        win.paste_selection()
        assert win.model.df.iat[0, 0] == "0x" and win.model.df.iat[4, 2] == "4z"
        assert win.model.df.iat[5, 0] == 6.0                     # 块外未动
        assert len(hits) == 1 and hits[0] == (1, 0, 5, 2)        # 一次外接矩形
        assert win.model.modified
        assert win.model.undo()                                  # 一次撤销
        pd.testing.assert_frame_equal(win.model.df, before, check_dtype=False)
        assert not win.model.undo() or win.model.df.equals(before)

    def test_paste_under_filter_writes_correct_source_rows(self, big):
        win = big
        win.active_filters = [{"col": "B", "condition": "大于", "value": "30"}]
        win._reapply_filters()
        assert len(win.model.df) == 3                            # 原表第 3、4、5 行
        win.table.setCurrentIndex(win.model.index(1, 0))
        QApplication.clipboard().setText("x\ny")
        win.paste_selection()
        assert list(win.model.df.iloc[:, 0].astype(str)) == ["x", "y", "6.0"]
        src = win.original_df.iloc[:, 0].astype(str).tolist()
        assert src == ["1.0", "2.0", "3.0", "x", "y", "6.0"]     # 只写回对应源行
        win.clear_all_filters()
        assert win.model.df.iloc[:, 0].astype(str).tolist() == src

    def test_fill_handle_down_column_undone_in_one_step(self, big):
        win = big
        win.model.setData(win.model.index(1, 1), "=A2*2")          # 源格公式
        before = win.model.df.copy()
        hits = _count_data_changes(win.model)
        win.table._fill_source = (1, 1, 1, 1)
        win.table._fill_target = (1, 1, 6, 1)
        win.table._perform_fill()
        win.table._fill_source = win.table._fill_target = None
        assert [win.model.formulas.get((r, 1)) for r in range(6)] == [
            "=A2*2", "=A3*2", "=A4*2", "=A5*2", "=A6*2", "=A7*2"]
        assert list(win.model.df["B"]) == [2.0, 4.0, 6.0, 8.0, 10.0, 12.0]
        assert len(hits) == 1
        assert tr("已填充 {} 个单元格").format(5) == win.statusBar().currentMessage()
        assert win.model.undo()                                  # 一步撤掉整次填充
        pd.testing.assert_frame_equal(win.model.df, before, check_dtype=False)
        assert win.model.formulas == {(0, 1): "=A2*2"}

    def test_ctrl_enter_fill_single_undo(self, big):
        win = big
        win.table.setCurrentIndex(win.model.index(1, 2))
        win.table.edit(win.model.index(1, 2))
        editor = win._cell_delegate.active_editor
        editor.setText("k")
        _select(win, 1, 2, 6, 2)
        win.table.commit_and_fill(editor)
        assert list(win.model.df["C"]) == ["k"] * 6
        assert win.model.undo()                                  # 撤掉批量填充（5 格）
        assert list(win.model.df["C"]) == ["k", "q", "r", "s", "t", "u"]
        assert win.model.undo()                                  # 再撤编辑器本身的提交
        assert list(win.model.df["C"]) == ["p", "q", "r", "s", "t", "u"]
