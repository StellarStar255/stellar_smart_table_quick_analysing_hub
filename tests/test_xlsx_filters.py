# -*- coding: utf-8 -*-
"""筛选写进 xlsx：Excel AutoFilter + 隐藏行 + 自定义属性里的精确副本；
打开时优先读文件（换电脑也在），Excel 里改过的筛选以 AutoFilter 为准。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
from PyQt6.QtCore import QSettings
from PyQt6.QtWidgets import QApplication

_app = QApplication.instance() or QApplication([])

from qtui import main_window as mw_mod
from qtui import xlsx_filters as xf
from qtui.main_window import MainWindow

BJ = {"col": "城市", "condition": "值在列表中", "value": ["北京"]}
BIG = {"col": "数量", "condition": "大于", "value": "1"}
HAS_A = {"col": "备注", "condition": "包含", "value": "a*b"}
NOT_EMPTY = {"col": "备注", "condition": "不为空", "value": ""}
ENDS = {"col": "备注", "condition": "结尾是", "value": "z"}


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    ini = str(tmp_path / "settings.ini")
    monkeypatch.setattr(mw_mod, "_make_settings",
                        lambda: QSettings(ini, QSettings.Format.IniFormat))


@pytest.fixture
def windows():
    made = []

    def make(path=None):
        w = MainWindow()
        made.append(w)
        if path:
            w.load_file(path)
            _app.processEvents()
        return w
    yield make
    for w in made:
        w.model.modified = False
        w.close()


def _book(path, frames):
    with pd.ExcelWriter(path) as xw:
        for name, df in frames.items():
            df.to_excel(xw, sheet_name=name, index=False)
    return path


@pytest.fixture
def xlsx(tmp_path):
    return _book(str(tmp_path / "b.xlsx"), {
        "甲": pd.DataFrame({"城市": ["北京", "上海", "北京", "广州"],
                           "数量": [1, 2, 3, 4],
                           "备注": ["xa*bz", "", "a*b", "q"]}),
        "乙": pd.DataFrame({"城市": ["广州", "北京"], "数量": [5, 6],
                           "备注": ["", ""]}),
    })


def _apply(w, *filters):
    w.active_filters = [dict(f) for f in filters]
    w._reapply_filters()


def _forget_local(monkeypatch, tmp_path):
    """模拟换了台电脑：本机侧车配置清空。"""
    monkeypatch.setattr(mw_mod, "FILE_CONFIG_PATH", str(tmp_path / "other_pc.json"))


class TestSpec:
    def test_mapping(self):
        spec = xf.autofilter_spec([BJ, BIG, HAS_A, ENDS], ["城市", "数量", "备注"])
        assert spec["0"] == {"values": ["北京"], "blank": False}
        assert spec["1"] == {"custom": [["greaterThan", "1"]], "and": False}
        assert spec["2"] == {"custom": [["equal", "*a~*b*"], ["equal", "*z"]], "and": True}

    def test_roundtrip_back_to_conditions(self):
        spec = xf.autofilter_spec([BJ, HAS_A, NOT_EMPTY], ["城市", "数量", "备注"])
        filters, skipped = xf.spec_to_filters(spec)
        assert skipped == 0
        assert xf.resolve_columns(filters, ["城市", "数量", "备注"]) == [
            BJ, HAS_A, NOT_EMPTY]

    def test_long_values_chunked_into_properties(self):
        wb = Workbook()
        many = {"col": "城市", "condition": "值在列表中",
                "value": ["城市{}".format(i) for i in range(300)]}
        xf.write_properties(wb, {"S": {"filters": [many], "sig": "x"}})
        assert all(len(p.value) <= 255 for p in wb.custom_doc_props)
        assert xf.read_properties_from_workbook(wb)["S"]["filters"] == [many]


class TestSaveAndOpen:
    def test_saved_into_file_and_visible_in_excel(self, windows, xlsx, monkeypatch, tmp_path):
        w = windows(xlsx)
        _apply(w, BJ)
        _apply(w, BJ, BIG)
        assert w.model.modified
        assert w.save_file() is True
        _app.processEvents()
        ws = load_workbook(xlsx)["甲"]
        assert ws.auto_filter.ref == "A1:C5"
        hidden = [r for r in range(2, 6) if ws.row_dimensions[r].hidden]
        assert hidden == [2, 3, 5]                  # 只剩 北京/3 那行
        _forget_local(monkeypatch, tmp_path)
        w2 = windows(xlsx)
        assert w2.active_filters == [BJ, BIG]
        assert w2.model.df["数量"].tolist() == [3]
        assert not w2.model.modified

    def test_exact_copy_keeps_what_autofilter_cannot(self, windows, xlsx, monkeypatch, tmp_path):
        w = windows(xlsx)
        _apply(w, HAS_A, NOT_EMPTY, ENDS)          # 同列 3 个自定义条件
        w.save_file()
        _app.processEvents()
        _forget_local(monkeypatch, tmp_path)
        w2 = windows(xlsx)
        assert w2.active_filters == [HAS_A, NOT_EMPTY, ENDS]
        assert w2.model.df["备注"].tolist() == ["xa*bz"]

    def test_other_sheet_filters_travel_too(self, windows, xlsx, monkeypatch, tmp_path):
        w = windows(xlsx)
        w.switch_sheet("乙")
        _apply(w, BIG)
        w.switch_sheet("甲")
        _apply(w, BJ)
        w.save_file()
        _app.processEvents()
        _forget_local(monkeypatch, tmp_path)
        w2 = windows(xlsx)
        assert w2.active_filters == [BJ]
        w2.switch_sheet("乙")
        assert w2.active_filters == [BIG]

    def test_clearing_removes_autofilter_and_unhides(self, windows, xlsx):
        w = windows(xlsx)
        _apply(w, BJ)
        w.save_file()
        _app.processEvents()
        w.clear_all_filters()
        w.save_file()
        _app.processEvents()
        ws = load_workbook(xlsx)["甲"]
        assert ws.auto_filter.ref is None
        assert not any(ws.row_dimensions[r].hidden for r in range(2, 6))
        assert xf.read_saved_filters(xlsx) is None

    def test_filter_set_in_excel_is_read(self, windows, tmp_path, monkeypatch):
        path = str(tmp_path / "excel.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"
        for row in (["城市", "数量"], ["北京", 1], ["上海", 2], ["北京", 3]):
            ws.append(row)
        ws.auto_filter = AutoFilter(ref="A1:B4", filterColumn=[
            FilterColumn(colId=0, filters=Filters(filter=["上海"]))])
        wb.save(path)
        w = windows(path)
        assert w.active_filters == [
            {"col": "城市", "condition": "值在列表中", "value": ["上海"]}]
        assert len(w.model.df) == 1

    def test_stale_exact_copy_loses_to_excel_change(self, windows, xlsx, monkeypatch, tmp_path):
        """别人在 Excel 里改了筛选（自定义属性被原样保留），以文件的 AutoFilter 为准。"""
        w = windows(xlsx)
        _apply(w, BJ)
        w.save_file()
        _app.processEvents()
        wb = load_workbook(xlsx)
        wb["甲"].auto_filter = AutoFilter(ref="A1:C5", filterColumn=[
            FilterColumn(colId=0, filters=Filters(filter=["广州"]))])
        wb.save(xlsx)
        _forget_local(monkeypatch, tmp_path)
        w2 = windows(xlsx)
        assert w2.active_filters == [
            {"col": "城市", "condition": "值在列表中", "value": ["广州"]}]

    def test_untouched_sheet_keeps_excel_autofilter(self, windows, tmp_path):
        """没在应用里动过筛选的 sheet，保存时原样保留文件里的 AutoFilter（含应用读不懂的）。"""
        path = str(tmp_path / "keep.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "A"
        for row in (["x"], [1], [2]):
            ws.append(row)
        ws.auto_filter.ref = "A1:A3"
        ws.row_dimensions[3].hidden = True        # 例如颜色筛选造成的隐藏
        wb.save(path)
        w = windows(path)
        w.model.setData(w.model.index(1, 0), "9")
        w.save_file()
        _app.processEvents()
        ws = load_workbook(path)["A"]
        assert ws.auto_filter.ref == "A1:A3"
        assert ws.row_dimensions[3].hidden

    def test_old_local_only_filters_get_written_on_save(self, windows, xlsx):
        """旧版只记在本机的 xlsx 筛选：打开恢复后，下次保存写进文件。"""
        w = windows(xlsx)
        _apply(w, BJ)
        w.model.modified = False
        w.close()                                  # 未保存：文件里没有，本机配置有
        assert xf.read_saved_filters(xlsx) is None
        w2 = windows(xlsx)
        assert w2.active_filters == [BJ]
        w2.model.setData(w2.model.index(1, 1), "7")
        w2.save_file()
        _app.processEvents()
        assert xf.read_saved_filters(xlsx) == {"甲": [BJ]}

    def test_csv_filter_change_not_modified(self, windows, tmp_path):
        path = str(tmp_path / "a.csv")
        pd.DataFrame({"城市": ["北京", "上海"], "数量": [1, 2]}).to_csv(path, index=False)
        w = windows(path)
        _apply(w, BJ)
        assert not w.model.modified


class TestLazyAndRename:
    def test_other_sheet_scanned_on_first_switch(self, windows, tmp_path, monkeypatch):
        path = str(tmp_path / "two.xlsx")
        wb = Workbook()
        a = wb.active
        a.title = "A"
        a.append(["x"]); a.append([1])
        b = wb.create_sheet("B")
        for row in (["城市"], ["北京"], ["上海"]):
            b.append(row)
        b.auto_filter = AutoFilter(ref="A1:A3", filterColumn=[
            FilterColumn(colId=0, filters=Filters(filter=["北京"]))])
        wb.save(path)
        calls = []
        real = xf.read_saved_filters
        monkeypatch.setattr(xf, "read_saved_filters",
                            lambda p, sheets=None: calls.append(sheets) or real(p, sheets))
        w = windows(path)
        assert calls == [["A"]]                    # 打开时只扫当前 sheet
        w.switch_sheet("B")
        assert w.active_filters == [
            {"col": "城市", "condition": "值在列表中", "value": ["北京"]}]
        assert len(w.model.df) == 1

    def test_renamed_sheet_keeps_filters_after_save(self, windows, xlsx, monkeypatch, tmp_path):
        w = windows(xlsx)
        w.switch_sheet("乙")
        _apply(w, BIG)
        w.switch_sheet("甲")
        w.save_file()
        _app.processEvents()
        w.model.modified = False
        w.close()
        _forget_local(monkeypatch, tmp_path)
        w2 = windows(xlsx)
        from PyQt6.QtWidgets import QInputDialog
        monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("丙", True))
        w2.rename_sheet("乙")                     # 没切过去就改名
        w2.save_file()
        _app.processEvents()
        w2.model.modified = False
        w2.close()
        _forget_local(monkeypatch, tmp_path / "x")
        assert xf.read_saved_filters(xlsx) is not None
        w3 = windows(xlsx)
        w3.switch_sheet("丙")
        assert w3.active_filters == [BIG]
